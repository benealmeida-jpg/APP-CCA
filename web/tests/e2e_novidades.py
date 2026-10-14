"""Conversa por cliente, quadro com processo parado e relatórios em Excel (modo demonstração)."""
import os, re, sys, subprocess, time, tempfile
from playwright.sync_api import sync_playwright, expect

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAIDA = sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp()
os.makedirs(SAIDA, exist_ok=True)
srv = subprocess.Popen([sys.executable, '-m', 'http.server', '8767', '--bind', '127.0.0.1'], cwd=RAIZ, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
time.sleep(1)
URL = 'http://127.0.0.1:8767/index.html?demo'
erros = []

def foto(pg, nome):
    pg.screenshot(path=os.path.join(SAIDA, f'{nome}.png'), full_page=True)

def entrar(pg, email):
    pg.goto(URL)
    pg.fill('#email', email); pg.fill('#senha', 'demo'); pg.click('button[type=submit]')

try:
    with sync_playwright() as p:
        b = p.chromium.launch(executable_path='/opt/pw-browsers/chromium-1194/chrome-linux/chrome')
        ctx = b.new_context(viewport={'width': 390, 'height': 844}, device_scale_factor=2, is_mobile=True, has_touch=True, accept_downloads=True)
        pg = ctx.new_page()
        pg.on('pageerror', lambda e: erros.append(f'pageerror: {e}'))
        pg.on('console', lambda m: erros.append(f'console: {m.text}') if m.type == 'error' and 'fonts.googleapis' not in m.text and 'ERR_' not in m.text else None)

        # 1) Corretor: aviso de mensagem nova na lista, conversa na ficha
        entrar(pg, 'corretor@demo')
        pg.wait_for_timeout(900); pg.evaluate("location.hash = '#/clientes'"); expect(pg.get_by_role('heading', name='Meus clientes')).to_be_visible()
        expect(pg.locator('tr', has_text='Mariana Costa').locator('.chip-msg')).to_contain_text('2 da CCA')
        expect(pg.get_by_role('button', name=re.compile(r'^Mensagens novas · 1'))).to_be_visible()
        foto(pg, 'n1-lista-corretor')
        pg.get_by_text('Mariana Costa').click()
        expect(pg.get_by_role('heading', name='Conversa com a CCA')).to_be_visible()
        expect(pg.get_by_text('A conta de luz veio com o endereço cortado', exact=False)).to_be_visible()
        expect(pg.get_by_text('Ebenézer · CCA').first).to_be_visible()
        pg.fill('#conversa-portal', 'Enviei a conta de luz nova agora há pouco.')
        pg.get_by_role('button', name='Enviar', exact=True).click()
        expect(pg.get_by_text('Mensagem enviada. A equipe da CCA foi avisada.')).to_be_visible()
        expect(pg.locator('.conversa-msg.minha', has_text='Enviei a conta de luz nova')).to_be_visible()
        pg.locator('#conversa').scroll_into_view_if_needed()
        foto(pg, 'n2-ficha-conversa')
        pg.get_by_role('link', name='Meus clientes').first.click()
        pg.wait_for_timeout(900); pg.evaluate("location.hash = '#/clientes'"); expect(pg.get_by_role('heading', name='Meus clientes')).to_be_visible()
        expect(pg.locator('tr', has_text='Mariana Costa').locator('.chip-msg')).to_have_count(0)
        pg.click('.usuario-botao'); pg.get_by_role('menuitem', name='Sair').click(); time.sleep(0.8)
        expect(pg.locator('#email')).to_be_visible()

        # 2) Equipe no computador: alertas, filtro Parados, quadro e conversa
        pg.set_viewport_size({'width': 1366, 'height': 900})
        entrar(pg, 'equipe@demo')
        pg.wait_for_timeout(900); pg.evaluate("location.hash = '#/painel/clientes'"); expect(pg.get_by_role('heading', name='Clientes', exact=True)).to_be_visible()
        expect(pg.get_by_text('Parado há 8 dias')).to_be_visible()
        expect(pg.locator('.chip-msg').first).to_be_visible()
        pg.get_by_role('button', name=re.compile(r'^Parados há 5\+ dias · 2')).click()
        expect(pg.locator('table.tabela tbody tr')).to_have_count(2)
        pg.get_by_role('button', name=re.compile(r'^Todos · ')).click()
        pg.get_by_role('button', name='Quadro', exact=True).click()
        expect(pg.locator('.quadro-coluna')).to_have_count(8)
        foto(pg, 'n3-quadro')
        pg.locator('.quadro-cartao', has_text='Mariana Costa').click()
        expect(pg.get_by_role('heading', name='Conversa com o corretor')).to_be_visible()
        expect(pg.locator('.conversa-msg', has_text='Enviei a conta de luz nova')).to_be_visible()
        pg.fill('#conversa-painel', 'Recebido, João. Vamos conferir hoje.')
        pg.get_by_role('button', name='Enviar', exact=True).click()
        expect(pg.get_by_text('Mensagem enviada. O corretor foi avisado.')).to_be_visible()
        foto(pg, 'n4-quadro-conversa')
        pg.get_by_role('button', name='Lista', exact=True).click()
        expect(pg.locator('table.tabela')).to_be_visible()

        # 3) Configurações: processo parado e lembretes
        pg.click('.menu-botoes a:has-text("Configurações")')
        expect(pg.get_by_role('heading', name='Processo parado e lembretes')).to_be_visible()
        pg.fill('#al-dias_alerta_parado', '10')
        pg.locator('#al-dias_alerta_parado').locator('xpath=..').get_by_role('button', name='Salvar').click()
        expect(pg.get_by_text('Salvo.')).to_be_visible()
        pg.click('.menu-botoes a:has-text("Clientes")')
        pg.wait_for_timeout(900); pg.evaluate("location.hash = '#/painel/clientes'"); expect(pg.get_by_role('heading', name='Clientes', exact=True)).to_be_visible()
        expect(pg.locator('table.tabela tbody tr', has_text='Mariana Costa')).to_be_visible()
        expect(pg.get_by_role('button', name=re.compile(r'^Parados há 10\+ dias'))).to_have_count(0)
        expect(pg.get_by_text('Parado há 8 dias')).to_have_count(0)

        # 4) Relatórios e Excel
        pg.click('.menu-botoes a:has-text("Relatórios")')
        expect(pg.get_by_role('heading', name='Relatórios')).to_be_visible()
        expect(pg.locator('.kpi', has_text='cadastros enviados')).to_contain_text('4')
        expect(pg.get_by_role('heading', name='Tempo em cada etapa')).to_be_visible()
        foto(pg, 'n5-relatorios')
        with pg.expect_download() as dl:
            pg.get_by_role('button', name='Baixar Excel').click()
        caminho = os.path.join(SAIDA, 'relatorio.xlsx')
        dl.value.save_as(caminho)
        import openpyxl
        wb = openpyxl.load_workbook(caminho)
        assert wb.sheetnames == ['Resumo', 'Clientes', 'Por status', 'Por imobiliária', 'Por corretor', 'Tempo por etapa'], wb.sheetnames
        assert wb['Clientes'].max_row == 5, wb['Clientes'].max_row
        print('excel ok:', wb.sheetnames)
        pg.get_by_role('button', name='Escolher datas').click()
        expect(pg.locator('#rel-de')).to_be_visible()
        pg.select_option('#rel-imob', label='Autônomo')
        expect(pg.locator('.kpi', has_text='cadastros enviados')).to_contain_text('1')

        # 5) Link do aviso (?c=) abre o cliente selecionado no painel
        mariana = pg.evaluate("JSON.parse(localStorage.getItem('appcca-demo-v9')).clientes.find(c => c.nome === 'Mariana Costa').id")
        pg.goto(f'http://127.0.0.1:8767/index.html?demo&c={mariana}')
        expect(pg.get_by_role('heading', name='Conversa com o corretor')).to_be_visible()
        expect(pg.locator('.lateral h2', has_text='Mariana Costa')).to_be_visible()
        assert pg.url.endswith('#/painel/clientes'), pg.url

        # 6) Celular: quadro e relatórios
        pg.set_viewport_size({'width': 390, 'height': 844})
        pg.get_by_role('button', name='Quadro', exact=True).click()
        foto(pg, 'n6-quadro-celular')
        pg.click('.menu-botoes a:has-text("Relatórios")')
        expect(pg.get_by_role('heading', name='Relatórios')).to_be_visible()
        largura = pg.evaluate('document.documentElement.scrollWidth')
        assert largura <= 392, f'rolagem lateral no celular: {largura}'
        foto(pg, 'n7-relatorios-celular')
        b.close()
finally:
    srv.terminate()

print('ERROS', erros)
assert not erros, erros
print('OK – conversa, quadro e relatórios')
